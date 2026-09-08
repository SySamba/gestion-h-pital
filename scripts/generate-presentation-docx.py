#!/usr/bin/env python3
"""Génère le document Word de présentation MedikaSN pour les hôpitaux."""

from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Installation requise: pip install python-docx")
    raise

OUTPUT = Path(__file__).resolve().parent.parent / "PRESENTATION_MEDIKASN_HOPITAL.docx"


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0B, 0x5E, 0x8A)
    return h


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def build():
    doc = Document()

    # Page de garde
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MedikaSN\n")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x0B, 0x5E, 0x8A)
    sub = title.add_run("Plateforme digitale de gestion hospitalière\nAdaptée au Sénégal")
    sub.font.size = Pt(16)

    doc.add_paragraph()
    p = doc.add_paragraph("Document de présentation commerciale")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True

    doc.add_page_break()

    add_heading(doc, "1. Résumé exécutif")
    doc.add_paragraph(
        "MedikaSN est une solution complète de gestion hospitalière conçue pour les cliniques, "
        "centres de santé et hôpitaux au Sénégal. Elle digitalise l'accueil, la file d'attente, "
        "les consultations, le laboratoire, la pharmacie et la direction — avec une interface "
        "web moderne et une application mobile pour les patients."
    )

    add_heading(doc, "2. Problématiques adressées", level=2)
    for item in [
        "Files d'attente longues et peu transparentes pour les patients",
        "Gestion manuelle des tickets et des salles de consultation",
        "Manque de visibilité pour la direction (statistiques, rapports)",
        "Communication limitée entre accueil, médecins et laboratoire",
        "Absence d'outil unifié adapté au contexte sénégalais (FCFA, Mobile Money, français)",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Modules fonctionnels", level=2)

    modules = [
        ("Accueil & réception", [
            "Enregistrement rapide des patients",
            "Consultation express (recherche + ticket en un flux)",
            "Gestion des guichets réception (configurable par l'admin)",
            "Émission de tickets avec paiement simulé (Mobile Money / Wave / Orange Money)",
        ]),
        ("File d'attente intelligente", [
            "Tickets numérotés avec QR code",
            "Durée de validité configurable par l'administrateur",
            "Expiration automatique si le patient ne se présente pas à temps",
            "Priorités : normale, urgente, très urgente",
            "Appel patient avec salle et médecin assignés",
        ]),
        ("Écran public salle d'attente", [
            "Affichage grand format : nom du patient, salle, ticket",
            "Annonce vocale en français (Edge TTS — voix naturelle)",
            "Répétition configurable (1 à 5 fois)",
            "Accès public via /ecran-attente (sans connexion)",
        ]),
        ("Alertes urgence / SOS", [
            "Bouton secours pour les patients en salle d'attente",
            "Notification immédiate au personnel (réception, médecins, admin)",
            "Suivi et traitement des alertes",
        ]),
        ("Consultations médicales", [
            "Dossier patient complet (allergies, antécédents, timeline)",
            "Historique médical et consultations",
            "Prescription d'ordonnances avec alertes allergies",
            "Export du dossier patient",
        ]),
        ("Rendez-vous", [
            "Prise de RDV en ligne par les patients",
            "Planning médecin avec disponibilités",
            "Rappels SMS simulés",
            "Vue du jour pour réception et médecins",
        ]),
        ("Laboratoire", [
            "Demandes d'analyses par les médecins",
            "Traitement par le laborantin",
            "Upload des résultats (PDF, images)",
            "Notification patient à la disponibilité des résultats",
        ]),
        ("Pharmacie", [
            "Consultation et délivrance des ordonnances",
            "Suivi du statut (active / délivrée)",
            "Sans module stock (focus sur le parcours patient)",
        ]),
        ("Administration complète", [
            "Paramètres système : nom hôpital, durée tickets, prix, temps d'attente",
            "CRUD salles / cabinets / box urgence",
            "CRUD guichets réception",
            "Création et gestion du personnel (médecins, réception, labo, pharmacie, admin)",
            "Activation / désactivation des comptes",
            "Multi-établissements (sites hospitaliers)",
        ]),
        ("Tableau de bord & rapports", [
            "Statistiques en temps réel : patients, RDV, file, tickets, recettes",
            "Activité récente et graphiques RDV sur 7 jours",
            "Rapports pour la direction",
        ]),
    ]

    for name, points in modules:
        add_heading(doc, name, level=3)
        for pt in points:
            add_bullet(doc, pt)

    add_heading(doc, "4. Rôles utilisateurs", level=2)
    roles = [
        ("Patient", "Tickets, file d'attente, RDV, analyses, ordonnances, dossier, SOS"),
        ("Réceptionniste", "Accueil, nouveaux patients, file, tickets, RDV, alertes urgence"),
        ("Médecin", "Patients, consultations, ordonnances, analyses, file d'attente"),
        ("Laborantin", "Analyses et upload des résultats"),
        ("Pharmacien", "Délivrance des ordonnances"),
        ("Administrateur", "Contrôle total : paramètres, salles, personnel, tous les modules"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Rôle"
    hdr[1].text = "Accès"
    for role, access in roles:
        row = table.add_row().cells
        row[0].text = role
        row[1].text = access

    doc.add_paragraph()

    add_heading(doc, "5. Avantages pour votre établissement", level=2)
    for item in [
        "Réduction du temps d'attente perçu grâce à la transparence (position, estimation)",
        "Meilleure organisation des salles et du personnel médical",
        "Image moderne avec écran d'appel vocal type hôpitaux internationaux",
        "Traçabilité complète du parcours patient",
        "Décisions éclairées via dashboard et rapports",
        "Solution prête pour le Sénégal : FCFA, numéros +221, interface en français",
        "Sécurité : authentification JWT, rôles et permissions stricts",
        "Déploiement flexible : serveur local ou cloud",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "6. Architecture technique", level=2)
    doc.add_paragraph(
        "Backend : Node.js / Express · Base de données MySQL · API REST sécurisée\n"
        "Frontend web : React / Vite · Interface responsive\n"
        "Mobile : React Native / Expo (patients)\n"
        "TTS : Microsoft Edge TTS (voix française naturelle, sans Docker requis)\n"
        "Uploads : résultats d'analyses (PDF, images)"
    )

    add_heading(doc, "7. Déploiement & accompagnement", level=2)
    for item in [
        "Installation sur votre infrastructure ou hébergement cloud",
        "Migration des données existantes (patients, personnel)",
        "Formation du personnel par rôle (1 à 2 jours)",
        "Personnalisation : logo, nom établissement, salles, guichets",
        "Support et maintenance évolutive",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "8. Démo live — scénario type", level=2)
    steps = [
        "Le patient achète un ticket (5 000 FCFA) via Mobile Money simulé",
        "Il suit sa position sur « Ma file d'attente » avec compte à rebours de validité",
        "La réception appelle le patient → écran public + annonce vocale ×3",
        "Le médecin consulte et prescrit une analyse + ordonnance",
        "Le laborantin upload le résultat PDF — le patient est notifié",
        "Le pharmacien délivre l'ordonnance",
        "L'admin consulte le dashboard et ajuste la durée des tickets",
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f"{i}. {step}")

    add_heading(doc, "9. Tarification indicative", level=2)
    doc.add_paragraph(
        "Nous proposons des formules adaptées à la taille de votre établissement :\n"
        "• Starter (clinique < 50 lits) — licence annuelle + formation\n"
        "• Pro (hôpital 50–200 lits) — multi-salles, écrans, support prioritaire\n"
        "• Enterprise (groupe hospitalier) — multi-sites, personnalisation avancée\n\n"
        "Contactez-nous pour un devis personnalisé incluant hébergement, formation et maintenance."
    )

    add_heading(doc, "10. Contact", level=2)
    contact = doc.add_paragraph()
    contact.add_run("MedikaSN — Santé digitale au Sénégal\n").bold = True
    contact.add_run("Email : contact@medikasn.sn\n")
    contact.add_run("Téléphone : +221 33 800 00 00\n")
    contact.add_run("Adresse : Dakar, Sénégal\n")
    contact.add_run("Site : www.medikasn.sn (à configurer)")

    doc.save(OUTPUT)
    print(f"Document créé : {OUTPUT}")


if __name__ == "__main__":
    build()
